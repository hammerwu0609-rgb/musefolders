"""Build the illustrated local user manual as an editable Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "使用手册.md"
OUTPUT = HERE / "曲境使用手册.docx"
FONT = "Microsoft YaHei"
GRAY = RGBColor(89, 96, 108)


def font(run, *, size: float | None = None, bold: bool | None = None,
         color: RGBColor | None = None) -> None:
    run.font.name = FONT
    if run._element.rPr is not None:
        rfonts = run._element.rPr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            run._element.rPr.insert(0, rfonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            font(run, size=9.6, color=GRAY)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            font(run, size=9, color=GRAY)
        else:
            run = paragraph.add_run(part)
            font(run)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tc_pr.append(edges)


def cell_padding(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", "95"), ("bottom", "95"),
                        ("left", "100"), ("right", "100")):
        side = OxmlElement(f"w:{name}")
        side.set(qn("w:w"), value)
        side.set(qn("w:type"), "dxa")
        margins.append(side)
    tc_pr.append(margins)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = ([1.05, 2.90, 2.90] if len(rows) > 4 else [1.28, 2.78, 2.79])
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            cell_padding(cell)
            if row_index == 0:
                shade(cell, "23354D")
            elif row_index % 2 == 0:
                shade(cell, "F4F7FA")
            p = cell.paragraphs[0]
            p.style = doc.styles["Normal"]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.14
            if col_index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                font(run, size=9, bold=row_index == 0 or (col_index == 0),
                     color=RGBColor(255, 255, 255) if row_index == 0 else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].startswith("|"):
        values = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def prepare_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.67)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(30, 35, 42)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(4)
    title = doc.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    title_ppr = title._element.get_or_add_pPr()
    for old_border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(old_border)
    title_border = OxmlElement("w:pBdr")
    for edge_name in ("top", "left", "bottom", "right", "between"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        title_border.append(edge)
    title_ppr.append(title_border)
    heading = doc.styles["Heading 1"]
    heading.font.name = FONT
    heading.font.size = Pt(13)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(11)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True


def add_image(doc: Document, image_markdown: str) -> None:
    match = re.match(r"!\[(.*?)\]\((.*?)\)", image_markdown)
    if not match:
        return
    alt, relative = match.groups()
    path = (HERE / relative).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(5.55))
    shape._inline.docPr.set("descr", alt)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    prepare_styles(doc)
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            inline(p, "曲境使用手册 本地演示版")
            index += 1
            continue
        if line.startswith("## "):
            heading = re.sub(r"^\d+\.\s*", "", line[3:])
            heading = (heading.replace("“① 找歌”", "找歌页面")
                       .replace("“② 审查”", "审查页面")
                       .replace("“③ 保存”", "保存页面"))
            p = doc.add_paragraph(style="Heading 1")
            inline(p, heading)
            index += 1
            continue
        if line.startswith("```text"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                index += 1
            index += 1
            flow = doc.add_paragraph()
            flow.paragraph_format.space_after = Pt(9)
            inline(flow, "选文件夹 → 分析 → 找歌 → 审查 → 保存")
            for run in flow.runs:
                font(run, size=10, bold=True, color=RGBColor(47, 75, 111))
            continue
        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if line.startswith("!["):
            add_image(doc, line)
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)", line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            inline(p, f"{numbered.group(1)}.  {numbered.group(2)}")
            index += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            inline(p, f"•  {line[2:]}")
            index += 1
            continue
        p = doc.add_paragraph()
        if line.startswith("*图："):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(7)
        inline(p, line)
        index += 1
    doc.core_properties.title = "曲境使用手册 本地演示版"
    doc.core_properties.subject = "曲境本地音乐地图的操作步骤"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
